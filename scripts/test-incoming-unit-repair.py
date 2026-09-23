"""Actual upload-id filtering tested on a disposable filesystem copy."""
import tempfile,shutil,json,subprocess
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
with tempfile.TemporaryDirectory() as temp:
    dest=Path(temp)
    names=['index.html','page-enhancements-core.js','task-unit-mappings.js','quality-file-routing.js','batch-policy.js','admin-auto-classifier.js','data/admin-import-index.json','data/north-quality-feedback-package.json','scripts/repair-quality-units.py','scripts/unit-evidence-context.js']
    for name in names:
        p=dest/name;p.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(ROOT/name,p)
    rel='data/质控意见反馈_管理员导入/test-incoming.docx';p=dest/rel;p.parent.mkdir(parents=True,exist_ok=True)
    d=Document();t=d.add_table(rows=1,cols=4)
    for c,text in zip(t.rows[0].cells,['县（市、区）','平山县','提交单位','中地科动察设计有限公司']):c.text=text
    d.add_paragraph('原文保留');d.save(p)
    index=dest/'data/admin-import-index.json';before=json.loads(index.read_text())
    entry={'kind':'quality-control','path':rel,'dataKey':'soilType','city':'石家庄市','district':'平山县','unit':'中地科勘察设计有限公司','batch':'2026年第二次第1批','uploadId':'fixture-new','complete':True,'unitCorrection':{'originalUnit':'中地科动察设计有限公司','unit':'中地科勘察设计有限公司'}}
    index.write_text(json.dumps(before+[entry],ensure_ascii=False))
    subprocess.run(['python',str(dest/'scripts/repair-quality-units.py'),'--mode','incoming','--upload-id','fixture-new'],check=True,cwd=dest)
    after=json.loads(index.read_text());assert after[:-1]==before,'Older index entries changed'
    assert Document(p).tables[0].cell(0,3).text=='中地科勘察设计有限公司'
    assert Document(p).paragraphs[0].text=='原文保留'
    audit=json.loads((dest/'data/unit-correction-history/fixture-new.json').read_text());assert audit['recordsChecked']==1 and audit['documentsChanged']==1,audit
    assert (dest/'index.html').read_bytes()==(ROOT/'index.html').read_bytes(),'Incoming job rewrote legacy data'
    assert after[-1]['unitCorrection']==entry['unitCorrection'] and after[-1]['directoryStatus']=='matched'
    print('Incoming repair passed: upload-id isolation, real DOCX repair, persistent evidence, original index and legacy data unchanged.')
