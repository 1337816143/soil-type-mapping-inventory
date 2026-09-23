"""Read-only mounted evidence plus disposable report fixtures; no network writes."""
import copy,hashlib,importlib.util,json,shutil,subprocess,tempfile
from pathlib import Path
import pymupdf as fitz
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
OLD='河北图宇地理信息科技有限公司';NEW='河北图宇科技有限公司'
spec=importlib.util.spec_from_file_location('tuyu_repair',ROOT/'scripts/repair-quality-units.py')
M=importlib.util.module_from_spec(spec);spec.loader.exec_module(M)
checks=[];snapshot=M.node()
rows=[r for r in snapshot['records'] if r.get('unitCorrection',{}).get('originalUnit')==OLD]
assert len(rows)==5 and {r['district'] for r in rows}=={'河间市','任丘市','肃宁县'}
receipt=json.loads((ROOT/'docs/tuyu-name-v1.2.10.json').read_text())
assert receipt['platformCorrections']==5 and receipt['documentsChanged']==0
assert receipt['directoryHashes']==snapshot['hashes']
for r in rows:
    assert r['dataKey']=='soilType' and r['city']=='沧州市' and r['unit']==NEW and not r['directory']['mismatch']
    item=next(x for x in receipt['reports'] if x['path']==r['path'])
    assert hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()==item['beforeSHA256']==item['afterSHA256']
    info=M.field(ROOT/r['path']);assert M.compact(info['value'])=='河北图宇' and M.qualified_geo(info,[r])
    checks.append('mounted '+r['district']+' '+r['batch'])
assert not any(r['unit']==OLD for r in snapshot['records'])
for city,district in [('沧州市','河间市'),('沧州市','任丘市'),('沧州市','肃宁县'),('石家庄市','无极县'),('石家庄市','深泽县'),('石家庄市','晋州市'),('石家庄市','藁城区')]:
    r=dict(dataKey='soilType',city=city,district=district,unit=OLD)
    answer=M.node([r])[0];assert answer['evidence']['action']=='typo' and answer['evidence']['unit']==NEW
    checks.append('typed exact rule '+city+district)
    for key in ['soilAttr','farmland','degradation','specialty','agricultural','landResource']:
        # Use real supported data keys below for the exhaustive negative set.
        if key not in snapshot['lists']:continue
        r['dataKey']=key;answer=M.node([r])[0]
        if NEW not in answer['directory']['listedUnits']:assert answer['evidence']['action']=='keep'
for key in snapshot['lists']:
    if key=='soilType':continue
    r=dict(dataKey=key,city='沧州市',district='河间市',unit=OLD)
    answer=M.node([r])[0];assert answer['evidence']['action']=='keep'
    checks.append('other result unchanged '+key)
for value in ['河北图宇','河北轩宇地理信息科技有限公司','河北图宇地理信息技术有限公司','河北图宇信息科技有限公司']:
    r=dict(dataKey='soilType',city='沧州市',district='河间市',unit=value)
    assert M.node([r])[0]['evidence']['action']=='keep';checks.append('not a fuzzy alias '+value)
for city,district in [('沧州市','青县'),('邯郸市','永年区'),('石家庄市','平山县')]:
    r=dict(dataKey='soilType',city=city,district=district,unit=OLD)
    assert M.node([r])[0]['evidence']['action']=='keep';checks.append('different assignment '+city+district)
TEMP=Path(tempfile.mkdtemp(prefix='_tuyu_fixture_',dir=ROOT/'data'))
def row(p,unit=OLD):return dict(path=str(p.relative_to(ROOT)),dataKey='soilType',city='沧州市',district='河间市',unit=unit,origin='legacy',batch='第一批')
try:
    path=TEMP/'wrong-header.pdf';d=fitz.open();p=d.new_page(width=595,height=842)
    font=fitz.Font('cjk');p.insert_font(fontname='fixture',fontbuffer=font.buffer)
    for x in [40,135,278,348,574]:p.draw_line((x,80),(x,120))
    for y in [80,120]:p.draw_line((40,y),(574,y))
    for x,t in [(46,'县（市、区）'),(141,'130984 河间市'),(284,'提交单位'),(354,OLD)]:p.insert_text((x,103),t,fontname='fixture',fontsize=9)
    p.insert_text((40,165),'正文保留：'+OLD,fontname='fixture',fontsize=10)
    p=d.new_page();p.insert_text((40,60),'Unchanged next page');d.save(path);d.close()
    r=row(path);a=M.process([r],ROOT/'test-artifacts/tuyu-pdf-fixture')[0]
    assert a['documentChanged'] and r['unit']==NEW and M.field(path)['value']==NEW
    with fitz.open(path) as result:assert OLD in result[0].get_text() and len(result)==2
    saved=path.read_bytes();assert not M.process([r])[0]['documentChanged'] and path.read_bytes()==saved
    checks.append('future PDF exact header-only edit and idempotence')
    path=TEMP/'wrong-header.docx';d=Document();t=d.add_table(rows=1,cols=4)
    for cell,value in zip(t.rows[0].cells,['县（市、区）','河间市','提交单位',OLD]):cell.text=value
    d.add_paragraph('正文保留：'+OLD);d.save(path)
    r=row(path);a=M.process([r])[0]
    assert a['documentChanged'] and Document(path).tables[0].cell(0,3).text==NEW and OLD in Document(path).paragraphs[0].text
    checks.append('future DOCX exact header only')
    path=TEMP/'manual.docx';d.save(path);r=row(path);r['unitSource']='manual';original=path.read_bytes();a=M.process([r])[0]
    assert r['unit']==OLD and not a['documentChanged'] and path.read_bytes()==original
    checks.append('manual choice retained')
    path=TEMP/'short.docx';t.cell(0,3).text='河北图宇';d.save(path);r=row(path);original=path.read_bytes();a=M.process([r])[0]
    assert r['unit']==NEW and not a['documentChanged'] and path.read_bytes()==original
    checks.append('genuine report shorthand preserved')
    assert snapshot['hashes']==M.node()['hashes'];checks.append('roster byte hashes unchanged')
finally:shutil.rmtree(TEMP)
OUT=ROOT/'test-artifacts';OUT.mkdir(exist_ok=True)
(OUT/'tuyu-name-tests.json').write_text(json.dumps({'status':'passed','count':len(checks),'checks':checks},ensure_ascii=False,indent=2))
print('TuYu evidence and regression checks passed:',len(checks))
